"""국회도서관(NANET) 「국회회의록 빅데이터」에서 내려받은 엑셀을 발언 단위 CSV로 바꾼다.
Convert a National Assembly Library minutes workbook (one sheet per meeting) into one row per utterance.

  python prepare_nanet.py 국회회의록.xlsx minutes.csv

* 내려받을 때 「회의록 내 모든 발언」 같은 옵션을 켜면 검색어와 무관한 발언이 섞인다. 받은 범위를 보고서 source에 적어 둔다.
* 이 엑셀은 시트 크기 정보가 비어 있는 경우가 있어 ws.reset_dimensions()가 없으면 첫 줄만 읽힌다.
* 발언자 칸은 대개 숫자 ID다. 이름은 infer_speaker_names.py로 추정하거나 공식 목록과 짝짓는다.
"""
import csv
import sys

import openpyxl


def main(src, out):
    wb = openpyxl.load_workbook(src, read_only=True)
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["meeting", "kind", "committee", "date", "speaker_id", "seq", "text"])
        for name in wb.sheetnames:
            ws = wb[name]
            ws.reset_dimensions()
            it = ws.iter_rows(values_only=True)
            hdr = None
            for r in it:
                if r and r[0] == "회의번호":
                    hdr = list(r)
                    break
            if not hdr or "발언자" not in hdr:
                continue
            ix = {h: i for i, h in enumerate(hdr)}
            tcols = sorted(i for h, i in ix.items() if h and str(h).startswith("발언내용"))
            for r in it:
                if not r or not r[0]:
                    continue
                text = "".join(str(r[i]) for i in tcols if i < len(r) and r[i]).strip()
                w.writerow([r[ix["회의번호"]], r[ix["회의구분"]], r[ix["위원회"]], str(r[ix["회의일자"]])[:11],
                            str(r[ix["발언자"]] or "").strip(), r[ix.get("발언순번", 0)], text])
                n += 1
    print(f"{n:,} utterances → {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
