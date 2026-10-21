# -*- coding: utf-8 -*-
"""수동 점검 — python tests/manual_check_inputs.py [이름 일부 …]
입력 자료 타입별 전 과정 점검: 넣기 → 설정 → 1차(실제 모델) → 판단 → 규칙 초안 → (일부) 2차.
화면이 쓰는 것과 같은 HTTP API로 돈다. 인터넷이 필요한 출처(유튜브·열린국회정보)는 이 PC 안의 가짜 서버로.
실제 임베딩 모델을 쓰므로 몇 분 걸린다. 미국 의회·Polis는 인터넷이 필요하다."""
import csv, io, json, os, re, sys, threading, time, urllib.request, urllib.error, tempfile, shutil
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORK = tempfile.mkdtemp(prefix="wwh-check-")
os.environ["WWH_WORK"] = WORK
sys.path.insert(0, REPO); sys.path.insert(0, REPO + "/servers")
import http_hub, sources  # noqa
import openpyxl

# ── 자료 만들기: 서로 다른 세 주제 · 두 집단 ──
TOP = {
 "ko": ["부동산 세금과 집값 대책을 더 강하게 해야 합니다", "검찰 개혁과 수사권 조정이 시급합니다", "청년 일자리와 최저임금 문제를 다뤄야 합니다"],
 "en": ["We need more buses and better transit on the east side", "Fix the potholes and repave the roads downtown", "Build more affordable housing for young families"],
}
def rows(lang="ko", n=60):
    out = []
    for g in ("A", "B"):
        for k, base in enumerate(TOP[lang]):
            for i in range(n):
                out.append({"id": f"{g}{k}{i}", "group": g, "author": f"u{(i*7+k)%23}", "text": f"{base} — {g} 의견 {i}번 {'정말' if i%2 else '꼭'} 그렇습니다" if lang=="ko" else f"{base}. Comment {g}-{i}, {'really' if i%2 else 'truly'} important."})
    return out

def csv_bytes(rs, delim=",", enc="utf-8"):
    b = io.StringIO(); w = csv.DictWriter(b, fieldnames=list(rs[0]), delimiter=delim); w.writeheader(); w.writerows(rs)
    return b.getvalue().encode(enc)

def xlsx_bytes(rs):
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(list(rs[0]))
    for r in rs: ws.append(list(r.values()))
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()

def nanet_xlsx(path, committee, kind, dates, sheets=2):
    wb = openpyxl.Workbook(); wb.remove(wb.active); k = 0
    for sh in range(sheets):
        ws = wb.create_sheet(f"회의{sh}")
        ws.append(["(국회도서관 발언 빅데이터)"])  # 머리 위 안내 줄
        ws.append(["회의번호", "회의구분", "위원회", "회의일자", "발언자", "발언순번", "발언내용1", "발언내용2"])
        for d in dates:
            for t, base in enumerate(TOP["ko"]):
                for i in range(12):
                    k += 1
                    ws.append([f"{d}-{sh}", kind, committee, d, "위원장", k, f"다음은 홍길동 위원님 질의해 주시기 바랍니다.", None])
                    k += 1
                    ws.append([f"{d}-{sh}", kind, committee, d, str(1000 + t), k, f"{base} {i}번째 질의입니다.", " 덧붙여 말씀드립니다."])
    wb.save(path)

# ── 가짜 외부 서버: 유튜브 API · 열린국회정보 · 일반 웹 ──
LIVE_CALLS = {"n": 0}
class Fake(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def send(self, code, body, ctype="application/json; charset=utf-8"):
        b = body if isinstance(body, bytes) else (body.encode() if isinstance(body, str) else json.dumps(body, ensure_ascii=False).encode())
        self.send_response(code); self.send_header("Content-Type", ctype); self.send_header("Content-Length", str(len(b))); self.end_headers(); self.wfile.write(b)
    def do_GET(self):
        u = urlparse(self.path); q = {k: v[0] for k, v in parse_qs(u.query).items()}; p = u.path
        if p.startswith("/yt/") and q.get("key") != "GOODKEY":
            return self.send(400, {"error": {"message": "API key not valid"}})
        if p == "/yt/videos":
            items = [{"id": v, "snippet": {"title": f"영상 {v}", "channelTitle": "채널갑" if v.startswith("a") else "채널을"},
                      "liveStreamingDetails": ({"activeLiveChatId": "CHAT1"} if v == "liveVideo01" else {})} for v in q["id"].split(",")]
            return self.send(200, {"items": items})
        if p == "/yt/commentThreads":
            vid = q["videoId"]; page = int(q.get("pageToken") or 0); items = []
            for i in range(50):
                n = page * 50 + i; base = TOP["ko"][n % 3]
                items.append({"snippet": {"topLevelComment": {"snippet": {"authorChannelId": {"value": f"UC{n%17}"}, "publishedAt": "2026-10-01T00:00:00Z", "likeCount": n % 5, "textDisplay": f"{base} ({vid} 댓글 {n})"}}},
                              "replies": {"comments": [{"snippet": {"authorChannelId": {"value": "UCr"}, "textDisplay": f"{base} 답글 {n}"}}]} if n % 10 == 0 else {}})
            return self.send(200, {"items": items, **({"nextPageToken": str(page + 1)} if page < 2 else {})})
        if p == "/yt/liveChat/messages":
            LIVE_CALLS["n"] += 1; c = LIVE_CALLS["n"]
            items = [{"snippet": {"type": "textMessageEvent", "publishedAt": "2026-10-01T00:00:00Z", "textMessageDetails": {"messageText": f"{TOP['ko'][(c*40+i)%3]} 채팅 {c}-{i}"}},
                      "authorDetails": {"channelId": f"UC{(c*40+i)%29}"}} for i in range(40)]
            body = {"items": items, "nextPageToken": f"t{c}", "pollingIntervalMillis": 100}
            if c >= 6: body["offlineAt"] = "2026-10-01T01:00:00Z"
            return self.send(200, body)
        if p.startswith("/oa/"):
            code = p.split("/")[-1]
            if q.get("KEY") != "OAKEY":
                return self.send(200, {"RESULT": {"CODE": "ERROR-290", "MESSAGE": "인증키가 유효하지 않습니다."}})
            page = int(q["pIndex"]); total = 230
            rows_ = [{"CONF_ID": f"C{page}-{i}", "COMM_NAME": "법제사법위원회", "CONF_TITLE": TOP["ko"][i % 3] + f" 회의 {page}-{i}", "PDF_LINK_URL": "http://x/y.pdf"}
                     for i in range(min(100, total - (page - 1) * 100))]
            return self.send(200, {code: [{"head": [{"list_total_count": total}, {"RESULT": {"CODE": "INFO-000"}}]}, {"row": rows_}]})
        if p == "/web/a.html":
            body = "<html><head><title>시민 토론</title></head><body><nav><p>메뉴 메뉴 메뉴 메뉴 메뉴 메뉴</p></nav>" + "".join(
                f"<p>{TOP['ko'][i%3]} 그래서 {i}번째 생각을 적습니다.</p>" for i in range(90)) + "<script>var x='<p>스크립트</p>'</script></body></html>"
            return self.send(200, body, "text/html; charset=utf-8")
        if p == "/web/b.html":
            return self.send(200, ("<html><head><meta charset='euc-kr'><title>공지</title></head><body>" + "".join(
                f"<li>{TOP['ko'][i%3]} 목록 {i}</li>" for i in range(60)) + "</body></html>").encode("cp949"), "text/html; charset=euc-kr")
        if p == "/data/d.csv":
            return self.send(200, csv_bytes(rows("en", 40)), "text/csv")
        if p == "/data/d.json":
            return self.send(200, {"data": rows("ko", 40)})
        if p == "/data/d.xlsx":
            return self.send(200, xlsx_bytes(rows("ko", 40)), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return self.send(404, {"error": "nf"})

fake = ThreadingHTTPServer(("127.0.0.1", 0), Fake); FP = fake.server_address[1]
threading.Thread(target=fake.serve_forever, daemon=True).start()
sources.YT = f"http://127.0.0.1:{FP}/yt/"
sources.OA_BASE = f"http://127.0.0.1:{FP}/oa/"

hub = http_hub.make_server(0); HP = hub.server_address[1]; B = f"http://127.0.0.1:{HP}"
threading.Thread(target=hub.serve_forever, daemon=True).start()
TOK = re.search(r'var TOKEN="([^"]+)"', urllib.request.urlopen(B + "/studio").read().decode()).group(1)

def api(m, path, body=None, raw=None, name=None):
    h = {"X-WWH-Token": TOK}
    data = None
    if raw is not None:
        data = raw; h["X-Filename"] = urllib.parse.quote(name)
    elif body is not None:
        data = json.dumps(body).encode()
    r = urllib.request.Request(B + "/api/" + path, data=data, method=m, headers=h)
    try:
        return json.load(urllib.request.urlopen(r, timeout=120))
    except urllib.error.HTTPError as e:
        return {"_http": e.code, "error": json.loads(e.read().decode() or "{}").get("error")}
import urllib.parse

def wait(jid, t=900):
    t0 = time.time()
    while time.time() - t0 < t:
        j = api("GET", "jobs/" + jid)
        if j.get("state") != "running":
            return j
        time.sleep(0.5)
    return {"state": "timeout"}

RESULTS = []
def record(name, ok, note):
    RESULTS.append((name, ok, note)); print(("OK  " if ok else "FAIL") + f" {name}: {note}", flush=True)

def analyse(pid, setup_over=None, round2=False):
    pr = api("GET", "p/" + pid)["project"]
    sug = dict(pr.get("suggest") or {})
    cols = pr["columns"]["columns"]
    st = {"text": sug.get("text") or ("text" if "text" in cols else cols[-1]), "group": sug.get("group") or ("group" if "group" in cols else cols[0]),
          "speaker": sug.get("speaker"), "context": sug.get("context") or [], "text_policy": sug.get("text_policy") or "private",
          "per_group": sug.get("per_group") or 0, "min_chars": sug.get("min_chars") or 10, "drop": sug.get("drop") or [],
          "selection": sug.get("selection") or "leaf", "min_cluster_size": sug.get("min_cluster_size") or 10,
          "embed_chars": sug.get("embed_chars"), "min_group_size": sug.get("min_group_size") or 0}
    st.update(setup_over or {})
    r = api("POST", f"p/{pid}/setup", st)
    if "_http" in r: return False, "setup " + str(r)
    j = wait(api("POST", f"p/{pid}/run", {"rules": None})["id"])
    if j["state"] != "done": return False, "round1 " + j.get("error", "")[:300]
    d = os.path.join(WORK, pid, "rounds")
    if not all(os.path.exists(os.path.join(d, f)) for f in ("r1_public.html", "r1_INTERNAL_do-not-share.html")):
        return False, "report files missing"
    pub = json.load(open(os.path.join(d, "r1_public.json"), encoding="utf-8"))
    rv = api("GET", f"p/{pid}/round/1/review")
    cl = rv["clusters"]
    # 원문 정책 점검
    texts_in_pub = sum(len(c.get("reps", [])) for c in pub["clusters"])
    authors_in_pub = sum(1 for c in pub["clusters"] for r in c.get("reps", []) if r.get("speaker")) + sum(1 for c in pub["clusters"] if c.get("top_speakers"))
    pol = st["text_policy"]
    pol_ok = (pol == "private" and texts_in_pub == 0 and authors_in_pub == 0) or (pol == "quotes" and texts_in_pub > 0 and authors_in_pub == 0) or (pol == "public" and texts_in_pub > 0)
    if not pol_ok: return False, f"text policy {pol}: texts {texts_in_pub}, authors {authors_in_pub}"
    more = api("GET", f"p/{pid}/round/1/cluster/{cl[0]['id']}?offset=0") if cl else {"texts": []}
    note = f"{pub['n_sample']}건 → 주제 {len(pub['clusters'])}개 (5판 {min(r['clusters'] for r in pub['runs'])}–{max(r['clusters'] for r in pub['runs'])}), 안 묶임 {pub['runs'][0]['unclustered_pct']}%, 공개판 원문 {texts_in_pub}·작성자 {authors_in_pub} ({pol}), 글 더 보기 {len(more['texts'])}/{more.get('total')}"
    if round2 and cl:
        api("POST", f"p/{pid}/round/1/review", {"reviewer": "A", "exclude": [cl[0]["id"]], "labels": {str(cl[1]['id']) if len(cl) > 1 else str(cl[0]['id']): "주제"}})
        api("POST", f"p/{pid}/round/1/review", {"reviewer": "B", "exclude": [cl[0]["id"]], "labels": {}})
        dr = api("POST", f"p/{pid}/round/1/draft", {"policy": "majority"})
        if "_http" in dr: return False, "draft " + str(dr)
        pats = dr["rules"]["exclude_patterns"] + [{"pattern": "(?i)테스트용패턴", "max_chars": 0}]
        sv = api("POST", f"p/{pid}/rules/r2.json", {"_from": dr["file"], "exclude_patterns": pats})
        if "_http" in sv: return False, "save rules " + str(sv)
        j2 = wait(api("POST", f"p/{pid}/run", {"rules": "r2.json"})["id"])
        if j2["state"] != "done": return False, "round2 " + j2.get("error", "")[:300]
        p2 = json.load(open(os.path.join(d, "r2_public.json"), encoding="utf-8"))
        lb = api("POST", f"p/{pid}/round/1/labels", {})
        note += f" | κ={dr['agreement']['kappa']} 초안 {len(dr['rules']['exclude_patterns'])}개 → 2차 주제 {len(p2['clusters'])}개, 차수기록 {len(p2['history'])}, 이름넣기 {'ok' if 'labels' in lb else lb}"
    return True, note

def project(name):
    return api("POST", "projects", {"name": name})["id"]

def fetch(pid, source, params):
    j = api("POST", f"p/{pid}/fetch", {"source": source, "params": params})
    if "_http" in j: return j
    return wait(j["id"], 600)

def case(name, fn):
    try:
        ok, note = fn()
    except Exception as e:
        import traceback; ok, note = False, f"{type(e).__name__}: {e} {traceback.format_exc()[-400:]}"
    record(name, ok, note)

T = tempfile.mkdtemp()
# ── 파일 ──
def c_csv():
    pid = project("csv-utf8"); r = api("POST", f"p/{pid}/upload", raw=csv_bytes(rows()), name="댓글.csv")
    return analyse(pid, {"text": "text", "group": "group", "speaker": "author"}, round2=True)
def c_cp949():
    pid = project("csv-cp949"); api("POST", f"p/{pid}/upload", raw=csv_bytes([dict(r, text=r["text"].replace("—", "-")) for r in rows()], enc="cp949"), name="엑셀저장.csv")
    return analyse(pid, {"text": "text", "group": "group", "text_policy": "quotes"})
def c_tsv():
    pid = project("tsv"); api("POST", f"p/{pid}/upload", raw=csv_bytes(rows("en"), delim="\t"), name="data.tsv")
    return analyse(pid, {"text": "text", "group": "group", "speaker": "author", "text_policy": "public"})
def c_xlsx():
    pid = project("xlsx"); api("POST", f"p/{pid}/upload", raw=xlsx_bytes(rows()), name="표.xlsx")
    return analyse(pid, {"text": "text", "group": "group"})
def c_path():
    p = os.path.join(T, "경로 자료.csv"); open(p, "wb").write(csv_bytes(rows()))
    pid = project("path"); r = api("POST", f"p/{pid}/path", {"path": p})
    if "_http" in r: return False, str(r)
    return analyse(pid, {"text": "text", "group": "group"})
def c_bad_upload():
    pid = project("bad"); r = api("POST", f"p/{pid}/upload", raw=b"MZ", name="x.exe")
    r2 = api("POST", f"p/{pid}/path", {"path": "/nope/none.csv"})
    return (r.get("_http") == 400 and r2.get("_http") == 404), f"exe → {r.get('_http')} {r.get('error')} / 없는 경로 → {r2.get('_http')}"

# ── 직접 받기 ──
def c_congress():
    pid = project("congress"); j = fetch(pid, "us-congress", {"start": "2026-03-03", "end": "2026-03-04"})
    if j.get("state") != "done": return False, str(j)[:300]
    return analyse(pid, {"per_group": 150, "min_group_size": 20})
def c_nanet():
    pid = project("nanet")
    f1 = os.path.join(T, "본회의.xlsx"); f2 = os.path.join(T, "법사위.xlsx")
    nanet_xlsx(f1, "본회의", "본회의", ["2025-03-04", "2025-09-10"]); nanet_xlsx(f2, "법제사법위원회", "상임위원회", ["2025-03-05"])
    names = [api("POST", f"p/{pid}/srcfile", raw=open(f, "rb").read(), name=os.path.basename(f))["file"] for f in (f1, f2)]
    j = fetch(pid, "korean-nanet", {"files": names, "date_from": "2025-03-01", "date_to": "2025-06-30"})
    if j.get("state") != "done": return False, str(j)[:300]
    pr = api("GET", "p/" + pid)["project"]
    ok, note = analyse(pid, {"min_chars": 5})
    return ok, f"받음 {pr['fetched']['rows']}행({pr['fetched']['note']}) | " + note
def c_open_assembly():
    pid = project("oa"); bad = fetch(pid, "open-assembly", {"key": "WRONG", "code": "nzmimeepazxkubdpn"})
    j = fetch(pid, "open-assembly", {"key": "OAKEY", "code": "nzmimeepazxkubdpn", "params": "DAE_NUM=22", "max_pages": 5})
    if j.get("state") != "done": return False, str(j)[:300]
    pr = api("GET", "p/" + pid)["project"]
    ok, note = analyse(pid, {"text": "CONF_TITLE", "group": "COMM_NAME", "min_chars": 5})
    return ok and bad.get("state") == "error", f"잘못된 키 → {bad.get('error','')[:60]} | 받음 {pr['fetched']['rows']}행 ({pr['fetched']['note'][:70]}) | " + note
def c_polis():
    pid = project("polis"); j = fetch(pid, "polis", {"conversation": "american-assembly.bowling-green"})
    if j.get("state") != "done": return False, str(j)[:300]
    return analyse(pid)
def c_yt_comments():
    pid = project("ytc"); bad = fetch(pid, "youtube-comments", {"key": "BAD", "videos": "https://youtu.be/aaaaaaaaaaa"})
    j = fetch(pid, "youtube-comments", {"key": "GOODKEY", "videos": "https://www.youtube.com/watch?v=aaaaaaaaaaa\nhttps://youtu.be/bbbbbbbbbbb, 진영B", "max_pages": 5})
    if j.get("state") != "done": return False, str(j)[:300]
    raw = open(os.path.join(WORK, pid, "input.csv"), encoding="utf-8").read()
    hashed = "UC1" not in raw
    keyfree = all("GOODKEY" not in open(os.path.join(r, f), encoding="utf-8", errors="ignore").read() for r, _, fs in os.walk(os.path.join(WORK, pid)) for f in fs if f.endswith((".json", ".csv")))
    ok, note = analyse(pid)
    return ok and hashed and keyfree and bad.get("state") == "error", f"잘못된 키 → 오류 · 작성자 해시 {hashed} · 키 저장 안 됨 {keyfree} | " + note
def c_yt_live():
    pid = project("ytl"); notlive = fetch(pid, "youtube-live", {"key": "GOODKEY", "url": "https://youtu.be/aaaaaaaaaaa"})
    j = fetch(pid, "youtube-live", {"key": "GOODKEY", "url": "https://www.youtube.com/live/liveVideo01?si=x", "max_minutes": 2})
    if j.get("state") != "done": return False, str(j)[:300]
    ok, note = analyse(pid, {"min_chars": 5})
    return ok and notlive.get("state") == "error", f"방송 아님 → {notlive.get('error','')[:50]} | {j.get('count')}건 받음 | " + note
def c_yt_live_stop():
    LIVE_CALLS["n"] = -1000  # 끝나지 않는 방송
    pid = project("ytl-stop"); r = api("POST", f"p/{pid}/fetch", {"source": "youtube-live", "params": {"key": "GOODKEY", "url": "liveVideo01", "max_minutes": 5}})
    time.sleep(2); api("POST", f"jobs/{r['id']}/stop", {}); j = wait(r["id"], 30)
    LIVE_CALLS["n"] = 0
    return j.get("state") == "done" and j.get("count", 0) > 0, f"멈추기 → {j.get('state')} · {j.get('count')}건에서 멈춤 ({j.get('elapsed')}s)"
def c_url_pages():
    pid = project("web"); j = fetch(pid, "url", {"urls": f"http://127.0.0.1:{FP}/web/a.html\nhttp://127.0.0.1:{FP}/web/b.html"})
    if j.get("state") != "done": return False, str(j)[:300]
    pr = api("GET", "p/" + pid)["project"]
    ok, note = analyse(pid, {"min_chars": 5})
    return ok, f"문단 {pr['fetched']['rows']}개(EUC-KR 페이지 포함) | " + note
def c_url_files():
    out = []
    for ext in ("csv", "json", "xlsx"):
        pid = project("url-" + ext); j = fetch(pid, "url", {"urls": f"http://127.0.0.1:{FP}/data/d.{ext}"})
        if j.get("state") != "done": return False, f"{ext}: {str(j)[:200]}"
        pr = api("GET", "p/" + pid)["project"]
        ok, note = analyse(pid, {"text": "text", "group": "group"})
        if not ok: return False, f"{ext}: {note}"
        out.append(f"{ext} {pr['columns']['rows']}행 → {note.split('→')[1].split(',')[0].strip()}")
    return True, " · ".join(out)

for name, fn in [("파일: CSV(UTF-8) + 2차까지", c_csv), ("파일: CSV(CP949·엑셀 저장)", c_cp949), ("파일: TSV(영어)", c_tsv),
                 ("파일: XLSX", c_xlsx), ("파일: 경로 입력", c_path), ("파일: 잘못된 입력 거절", c_bad_upload),
                 ("받기: 미국 의회 회의록", c_congress), ("받기: 국회도서관 XLSX 2개+기간 거르기", c_nanet),
                 ("받기: 열린국회정보 Open API(가짜 서버)", c_open_assembly), ("받기: Polis", c_polis),
                 ("받기: 유튜브 댓글(가짜 API)", c_yt_comments), ("받기: 유튜브 실시간 채팅(가짜 API)", c_yt_live),
                 ("받기: 실시간 채팅 멈추기", c_yt_live_stop), ("받기: 웹 페이지 2개", c_url_pages),
                 ("받기: URL 파일(CSV·JSON·XLSX)", c_url_files)]:
    if len(sys.argv) > 1 and not any(a in name for a in sys.argv[1:]):
        continue
    case(name, fn)

print("\n=== 요약 ===")
for n, ok, note in RESULTS:
    print(("✅" if ok else "❌"), n)
hub.shutdown(); fake.shutdown()
shutil.rmtree(WORK, ignore_errors=True); shutil.rmtree(T, ignore_errors=True)
