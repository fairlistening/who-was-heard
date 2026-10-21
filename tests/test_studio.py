"""분석 화면 뒷단 — 자료 → 설정 → 1차 → 판단 2명 → 규칙 초안 → 2차. 무거운 모델 없이 (가짜 임베딩).
Studio back end, end to end, with stand-in embedding/layout/clustering."""
import csv
import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "servers"))
sys.path.insert(0, os.path.join(ROOT, "tests"))

import whowasheard.report as R  # noqa: E402
from whowasheard.cluster import ClusterParams  # noqa: E402
import studio  # noqa: E402
from test_core import corpus, fake_cluster, fake_embed, fake_layout  # noqa: E402

_real = R.diagnose


def _fake_diagnose(*a, **kw):
    kw.update(embed_fn=fake_embed, layout_fn=fake_layout, cluster_fn=fake_cluster, cache_dir=None)
    p = kw.get("params") or ClusterParams()
    kw["params"] = ClusterParams(seeds=(42, 1, 7), min_cluster_size=p.min_cluster_size, selection=p.selection)
    return _real(*a, **kw)


def wait(job):
    for _ in range(300):
        j = studio.job_status(job["id"])
        if j["state"] != "running":
            return j
        time.sleep(0.05)
    raise AssertionError("job did not finish")


class Studio(unittest.TestCase):
    def setUp(self):
        self.work = tempfile.mkdtemp()
        os.environ["WWH_WORK"] = self.work
        R.diagnose = _fake_diagnose

    def tearDown(self):
        R.diagnose = _real
        os.environ.pop("WWH_WORK", None)
        shutil.rmtree(self.work, ignore_errors=True)

    def csv_bytes(self):
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=["txt", "grp", "who"])
        w.writeheader()
        w.writerows(corpus())
        return buf.getvalue().encode("utf-8")

    def test_full_flow(self):
        pid = studio.create_project("시연 Demo")["id"]
        r = studio.set_input_upload(pid, "data.csv", self.csv_bytes())
        cols = r["project"]["columns"]
        self.assertEqual(cols["columns"], ["txt", "grp", "who"])
        self.assertIn("grp", cols["values"])          # 값이 적은 열은 값을 보여 줌
        self.assertNotIn("txt", cols["values"])       # 글 열은 값을 보여 주지 않음
        studio.save_setup(pid, {"text": "txt", "group": "grp", "speaker": "who", "text_policy": "private", "per_group": 0})

        j = wait(studio.start_round(pid))
        self.assertEqual(j["state"], "done", j.get("log"))
        rv = studio.review_data(pid, 1)
        self.assertTrue(rv["clusters"])
        self.assertTrue(rv["clusters"][0]["samples"][0]["text"])   # 이 PC 화면에는 원문 표본
        ids = [c["id"] for c in rv["clusters"]]

        studio.save_review(pid, 1, {"reviewer": "A", "exclude": [ids[0]], "labels": {str(ids[1]): "학교"}})
        rv = studio.save_review(pid, 1, {"reviewer": "B", "exclude": [ids[0]], "labels": {str(ids[1]): "교육"}})
        self.assertEqual(len(rv["agreement"]["reviewers"]), 2)
        self.assertEqual(rv["agreement"]["excluded"], [ids[0]])

        d = studio.make_draft(pid, 1, "majority")
        self.assertEqual(d["file"], "r2-draft.json")
        saved = studio.save_rules(pid, "r2.json", {"_from": d["file"], "exclude_patterns":
                                                   d["rules"]["exclude_patterns"] + [{"pattern": "(?i)yield back", "max_chars": 300}]})
        self.assertIn({"pattern": "(?i)yield back", "max_chars": 300}, saved["rules"]["exclude_patterns"])

        j = wait(studio.start_round(pid, "r2.json"))
        self.assertEqual(j["state"], "done", j.get("log"))
        proj = studio.get_project(pid)["project"]
        self.assertEqual([r["n"] for r in proj["rounds"]], [1, 2])
        with open(os.path.join(self.work, pid, "rounds", "r2_public.json"), encoding="utf-8") as f:
            pub = json.load(f)
        self.assertEqual(pub["round"], 2)
        self.assertEqual(len(pub["history"]), 1)
        self.assertNotIn("comment North", json.dumps(pub, ensure_ascii=False))   # 공개판엔 원문 없음

        studio.apply_review_labels(pid, 1)
        with open(os.path.join(self.work, pid, "rounds", "r1_public.json"), encoding="utf-8") as f:
            lab = {c["id"]: c["label"] for c in json.load(f)["clusters"]}
        self.assertEqual(lab[ids[1]], "학교 / 교육")                            # 엇갈린 이름은 둘 다

    def test_guards(self):
        pid = studio.create_project("g")["id"]
        with self.assertRaises(ValueError):
            studio.save_rules(pid, "r2.json", {"exclude_patterns": [{"pattern": "(unclosed"}]})
        with self.assertRaises(ValueError):
            studio.save_rules(pid, "../x.json", {})
        with self.assertRaises(FileNotFoundError):
            studio.file_path(pid, "../../etc/passwd")
        with self.assertRaises(ValueError):
            studio._pdir("../x")
        with self.assertRaises(ValueError):
            studio.set_input_upload(pid, "x.exe", b"MZ")
        code, _, _ = studio.route("GET", "/api/projects", {}, b"", {"X-WWH-Token": "wrong"}, "right")
        self.assertEqual(code, 403)
        code, _, _ = studio.route("GET", f"/files/{pid}/project.json", {"t": ["wrong"]}, b"", {}, "right")
        self.assertEqual(code, 403)
        code, body, _ = studio.route("GET", "/api/projects", {}, b"", {"X-WWH-Token": "right"}, "right")
        self.assertEqual(code, 200)
        self.assertEqual(body["projects"][0]["id"], pid)


class Fetch(unittest.TestCase):
    """직접 받기 — 인터넷 없이: 이 PC 안의 작은 웹 서버와 가짜 국회도서관 엑셀로."""
    def setUp(self):
        self.work = tempfile.mkdtemp()
        os.environ["WWH_WORK"] = self.work

    def tearDown(self):
        os.environ.pop("WWH_WORK", None)
        shutil.rmtree(self.work, ignore_errors=True)

    def test_url_pages(self):
        import http.server
        import threading
        site = tempfile.mkdtemp()
        with open(os.path.join(site, "a.html"), "w", encoding="utf-8") as f:
            f.write("<html><head><title>Town hall</title></head><body><nav><p>menu menu menu menu menu menu</p></nav>"
                    "<p>We need more buses on the east side, the commute takes far too long.</p>"
                    "<p>short</p><ul><li>Fix the potholes on Main Street before winter comes.</li></ul>"
                    "<script>var x='<p>not this</p>'</script></body></html>")
        h = http.server.ThreadingHTTPServer(("127.0.0.1", 0), lambda *a: http.server.SimpleHTTPRequestHandler(*a, directory=site))
        threading.Thread(target=h.serve_forever, daemon=True).start()
        try:
            pid = studio.create_project("url")["id"]
            j = wait(studio.start_fetch(pid, "url", {"urls": f"http://127.0.0.1:{h.server_address[1]}/a.html"}))
            self.assertEqual(j["state"], "done", j.get("log"))
            proj = studio.get_project(pid)["project"]
            self.assertEqual(proj["columns"]["rows"], 3)          # 메뉴·스크립트는 빠짐 (5자 이상 문단은 남음)
            self.assertEqual(proj["suggest"]["text"], "text")
            self.assertEqual(proj["fetched"]["source"], "url")
        finally:
            h.shutdown()
            h.server_close()
            shutil.rmtree(site, ignore_errors=True)

    def test_nanet_workbook_and_no_key_saved(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["회의번호", "회의구분", "위원회", "회의일자", "발언자", "발언순번", "발언내용1"])
        k = 0
        for i in range(4):
            k += 1
            ws.append([1, "상임위", "법제사법위원회", "2026-03-03", "999", k, "다음은 홍길동 위원님 질의해 주시기 바랍니다."])
            k += 1
            ws.append([1, "상임위", "법제사법위원회", "2026-03-03", "1234", k, f"검찰 개혁에 대해 묻겠습니다 {i}"])
        src = os.path.join(self.work, "n.xlsx")
        wb.save(src)
        pid = studio.create_project("kn")["id"]
        with open(src, "rb") as f:
            name = studio.save_source_file(pid, "회의록.xlsx", f.read())["file"]
        j = wait(studio.start_fetch(pid, "korean-nanet", {"file": name}))
        self.assertEqual(j["state"], "done", j.get("log"))
        from whowasheard.data import load_table
        rows = load_table(os.path.join(self.work, pid, "input.csv"))
        self.assertIn("홍길동", {r["speaker"] for r in rows})
        # 키는 source.json에 남지 않는다
        pid2 = studio.create_project("yt")["id"]
        j = wait(studio.start_fetch(pid2, "youtube-live", {"key": "SECRET-KEY-123", "url": "not a url"}))
        self.assertEqual(j["state"], "error")
        for root, _, files in os.walk(self.work):
            for fn in files:
                if fn.endswith((".json", ".csv")):
                    with open(os.path.join(root, fn), encoding="utf-8", errors="ignore") as f:
                        self.assertNotIn("SECRET-KEY-123", f.read())

    def test_nanet_named_speakers_and_dates(self):
        import sources
        self.assertEqual(sources.iso_date("2025년7월28일"), "2025-07-28")
        self.assertEqual(sources.iso_date("2025년01월06일"), "2025-01-06")
        self.assertEqual(sources.split_speaker("서영교 위원"), ("서영교", "위원", "의원"))
        self.assertEqual(sources.split_speaker("소위원장 박수영"), ("박수영", "소위원장", "의원"))
        self.assertEqual(sources.split_speaker("금융위원회부위원장 김소영")[2], "정부·기관·증인 등")
        self.assertEqual(sources.split_speaker("수석전문위원 홍길동")[2], "국회 사무처")
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["회의번호", "회의구분", "위원회", "회의일자", "발언자", "발언순번", "발언내용1"])
        for i, (who, d) in enumerate([("서영교 위원", "2025년01월06일"), ("법무부차관 김석우", "2025년7월28일"), ("소위원장 김용민", "2025년12월1일")]):
            ws.append([1, "소위원회", "법제사법위원회-법안심사제1소위원회", d, who, i, f"검찰 개혁에 대해 말씀드립니다 {i}"])
        src = os.path.join(self.work, "n2.xlsx")
        wb.save(src)
        pid = studio.create_project("kn2")["id"]
        with open(src, "rb") as f:
            name = studio.save_source_file(pid, "소위.xlsx", f.read())["file"]
        j = wait(studio.start_fetch(pid, "korean-nanet", {"files": [name], "date_from": "2025-02-01", "date_to": "2025-12-31"}))
        self.assertEqual(j["state"], "done", j.get("log"))
        from whowasheard.data import load_table
        rows = load_table(os.path.join(self.work, pid, "input.csv"))
        self.assertEqual([r["date"] for r in rows], ["2025-07-28", "2025-12-01"])   # 1월은 기간 밖
        self.assertEqual([r["side"] for r in rows], ["정부·기관·증인 등", "의원"])
        self.assertEqual(rows[1]["speaker_name"], "김용민")

    def test_nanet_uploaded_as_plain_file_is_converted(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "목록"
        ws.append(["회의록 목록리스트"])
        ws.append(["054665"])
        ws2 = wb.create_sheet("054665")
        ws2.append(["회의번호", "회의구분", "위원회", "회의일자", "발언자", "발언순번", "발언내용1"])
        for i in range(4):
            ws2.append(["054665", "소위원회", "법제사법위원회-법안심사제1소위원회", "2025년01월06일", "서영교 위원", i,
                        f"생계 통장 압류를 금지하는 법안을 빨리 처리해야 합니다 {i}"])
        bio = io.BytesIO()
        wb.save(bio)
        pid = studio.create_project("kn3")["id"]
        r = studio.set_input_upload(pid, "제22대+소위원회.xlsx", bio.getvalue())
        self.assertEqual(r.get("redirected"), "korean-nanet")
        j = wait(r["job"])
        self.assertEqual(j["state"], "done", j.get("log"))
        cols = studio.get_project(pid)["project"]["columns"]["columns"]
        self.assertIn("text", cols)
        self.assertIn("side", cols)

    def test_empty_after_cleaning_message(self):
        from whowasheard.report import diagnose
        from whowasheard.clean import CleanRules
        with self.assertRaises(ValueError) as cm:
            diagnose([{"t": "목록", "g": "a"}] * 5, "t", "g", rules=CleanRules(min_chars=30))
        self.assertIn("짧음 5", str(cm.exception))
        self.assertIn("평균 길이는 2자", str(cm.exception))

    def test_video_ids(self):
        import sources
        for u in ("https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=1", "https://youtu.be/dQw4w9WgXcQ",
                  "youtube.com/live/dQw4w9WgXcQ?si=x", "dQw4w9WgXcQ"):
            self.assertEqual(sources.video_id(u), "dQw4w9WgXcQ", u)
        self.assertIsNone(sources.video_id("https://example.com/"))


if __name__ == "__main__":
    unittest.main()
